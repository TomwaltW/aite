"""aite_label 的回归测试（镜像里：python -m unittest discover -s /opt/aite/tests）。

夹具全部在 tempfile 里现造：pillow 造图、python-docx / openpyxl 造文档、PPTX 手搓最小 OPC zip
（镜像里没有 python-pptx，也不许加）。不联网、不靠 sleep。
"""

import json
import os
import re
import struct
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import aite_label  # noqa: E402

LABELER = os.path.join(ROOT, "aite_label.py")
FIVE = {"Label", "ContentProducer", "ProduceID", "ContentPropagator", "PropagateID"}

try:
    import pypdf  # noqa: F401

    HAVE_PYPDF = True
except ImportError:
    HAVE_PYPDF = False


def _fields(n="1"):
    return {
        "Label": "1",
        "ContentProducer": "Aite",
        "ProduceID": "produce-" + n,
        "ContentPropagator": "飞书",
        "PropagateID": "msg-" + n,
    }


def _png_text_chunks(data):
    """[(类型, keyword)]，只收 tEXt / iTXt / zTXt。"""
    out, pos = [], 8
    while pos + 8 <= len(data):
        (length,) = struct.unpack(">I", data[pos : pos + 4])
        ctype = data[pos + 4 : pos + 8]
        body = data[pos + 8 : pos + 8 + length]
        if ctype in (b"tEXt", b"iTXt", b"zTXt"):
            out.append((ctype.decode(), body.split(b"\x00", 1)[0].decode("latin-1")))
        pos += 12 + length
    return out


def _jpeg_scan(data):
    i = data.index(b"\xff\xda")
    return data[i:]


def _pixels(path):
    from PIL import Image

    with Image.open(path) as im:
        return im.mode, im.size, im.tobytes()


class LabelTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = self._tmp.name

    def tearDown(self):
        self._tmp.cleanup()

    def path(self, name):
        return os.path.join(self.dir, name)

    def make_png(self, name="a.png"):
        from PIL import Image, PngImagePlugin

        p = self.path(name)
        im = Image.new("RGB", (64, 48), (30, 120, 200))
        info = PngImagePlugin.PngInfo()
        info.add_text("Software", "matplotlib")
        im.save(p, pnginfo=info)
        return p

    def make_jpeg(self, name="a.jpg"):
        from PIL import Image

        p = self.path(name)
        im = Image.new("RGB", (64, 48), (200, 80, 40))
        exif = Image.Exif()
        exif[0x010F] = "TestCam"
        im.save(p, format="JPEG", quality=85, exif=exif.tobytes())
        return p

    # ------------------------------------------------------------ PNG

    def test_png_itxt_roundtrip(self):
        p = self.make_png()
        res = aite_label.label_file(p, _fields())
        self.assertEqual(res["status"], "labeled", res)
        self.assertEqual(res["format"], "png")
        with open(p, "rb") as f:
            chunks = _png_text_chunks(f.read())
        self.assertIn(("iTXt", "AIGC"), chunks)
        self.assertIn("Software", [k for _, k in chunks])
        self.assertEqual(aite_label.read_label(p), _fields())

    def test_png_relabel_replaces_not_duplicates(self):
        p = self.make_png()
        aite_label.label_file(p, _fields("1"))
        aite_label.label_file(p, _fields("2"))
        with open(p, "rb") as f:
            keys = [k for _, k in _png_text_chunks(f.read())]
        self.assertEqual(keys.count("AIGC"), 1, keys)
        self.assertEqual(aite_label.read_label(p)["ProduceID"], "produce-2")

    # ------------------------------------------------------------ JPEG

    def test_jpeg_exif_and_xmp_roundtrip(self):
        from PIL import Image

        p = self.make_jpeg()
        with open(p, "rb") as f:
            scan_before = _jpeg_scan(f.read())
        res = aite_label.label_file(p, _fields())
        self.assertEqual(res["status"], "labeled", res)
        with open(p, "rb") as f:
            data = f.read()
        # 扫描数据逐字节不变 = 没重压缩像素
        self.assertEqual(_jpeg_scan(data), scan_before)
        self.assertIn(aite_label.XMP_NS.encode(), data)
        with Image.open(p) as im:
            exif = im.getexif()
            self.assertEqual(exif.get(0x010F), "TestCam")
            comment = exif.get_ifd(0x8769).get(0x9286)
            xmp = im.info.get("xmp")
        self.assertTrue(comment.startswith(b"ASCII\x00\x00\x00"), comment)
        self.assertEqual(json.loads(comment[8:].decode("ascii")), _fields())
        self.assertIn(aite_label.XMP_NS.encode(), xmp)
        self.assertEqual(aite_label.read_label(p), _fields())
        aite_label.label_file(p, _fields("2"))
        with open(p, "rb") as f:
            self.assertEqual(f.read().count(aite_label.XMP_NS.encode()), 1)
        self.assertEqual(aite_label.read_label(p)["ProduceID"], "produce-2")

    # ------------------------------------------------------------ OOXML

    def test_docx_custom_props_roundtrip(self):
        import docx

        p = self.path("a.docx")
        d = docx.Document()
        d.add_paragraph("季度报告")
        d.save(p)
        res = aite_label.label_file(p, _fields())
        self.assertEqual(res["status"], "labeled", res)
        self.assertEqual(aite_label.read_label(p), _fields())
        with zipfile.ZipFile(p) as z:
            self.assertIn("docProps/custom.xml", z.namelist())
            self.assertIn(b'name="AIGC"', z.read("docProps/custom.xml"))
        self.assertEqual(docx.Document(p).paragraphs[0].text, "季度报告")

    def test_xlsx_keeps_existing_custom_props(self):
        import openpyxl
        from openpyxl.packaging.custom import StringProperty

        p = self.path("a.xlsx")
        wb = openpyxl.Workbook()
        wb.active["A1"] = "销量"
        wb.custom_doc_props.append(StringProperty(name="Owner", value="张三"))
        wb.save(p)
        res = aite_label.label_file(p, _fields())
        self.assertEqual(res["status"], "labeled", res)
        with zipfile.ZipFile(p) as z:
            xml = z.read("docProps/custom.xml").decode("utf-8")
        root = ET.fromstring(xml.encode("utf-8"))
        by_name = {p.get("name"): int(p.get("pid")) for p in root}
        self.assertIn("Owner", by_name, xml)
        self.assertIn("AIGC", by_name, xml)
        self.assertGreater(by_name["AIGC"], by_name["Owner"])
        wb2 = openpyxl.load_workbook(p)
        self.assertEqual(wb2.active["A1"].value, "销量")
        names = {prop.name: prop.value for prop in wb2.custom_doc_props}
        self.assertEqual(names.get("Owner"), "张三")
        self.assertEqual(json.loads(names["AIGC"]), _fields())

    def test_pptx_minimal_package_labelled(self):
        p = self.path("a.pptx")
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr(
                "[Content_Types].xml",
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Default Extension="xml" ContentType="application/xml"/>'
                '<Override PartName="/ppt/presentation.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
                "</Types>",
            )
            z.writestr(
                "_rels/.rels",
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                'Target="ppt/presentation.xml"/>'
                "</Relationships>",
            )
            z.writestr(
                "ppt/presentation.xml",
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>',
            )
        for n in ("1", "2"):
            res = aite_label.label_file(p, _fields(n))
            self.assertEqual(res["status"], "labeled", res)
        with zipfile.ZipFile(p) as z:
            ct = z.read("[Content_Types].xml").decode("utf-8")
            rels = z.read("_rels/.rels").decode("utf-8")
            self.assertIn("ppt/presentation.xml", z.namelist())
        override = re.findall(r'<Override PartName="/docProps/custom.xml" ContentType="([^"]+)"', ct)
        self.assertEqual(override, [aite_label._CUSTOM_CT], ct)
        custom_rel = re.findall(r"<Relationship [^>]*custom-properties[^>]*/>", rels)
        self.assertEqual(len(custom_rel), 1, rels)
        self.assertIn('Target="docProps/custom.xml"', custom_rel[0])
        self.assertIn('Id="rId2"', custom_rel[0])
        self.assertEqual(aite_label.read_label(p), _fields("2"))

    # ------------------------------------------------------------ Markdown

    def test_markdown_front_matter_added(self):
        p = self.path("a.md")
        with open(p, "w", encoding="utf-8") as f:
            f.write("# 周报\n\n正文里有 'quote'。\n")
        res = aite_label.label_file(p, _fields())
        self.assertEqual(res["status"], "labeled", res)
        with open(p, encoding="utf-8") as f:
            text = f.read()
        self.assertTrue(text.startswith("---\nAIGC: '"), text)
        self.assertTrue(text.endswith("---\n# 周报\n\n正文里有 'quote'。\n"), text)
        self.assertEqual(aite_label.read_label(p), _fields())

    def test_markdown_existing_front_matter_kept(self):
        p = self.path("a.markdown")
        with open(p, "w", encoding="utf-8") as f:
            f.write("---\ntitle: 报告\nAIGC: 'old'\ntags: [a, b]\n---\n# 正文\n")
        f2 = _fields()
        f2["ContentPropagator"] = "O'Brien"
        aite_label.label_file(p, f2)
        with open(p, encoding="utf-8") as f:
            text = f.read()
        self.assertEqual(text.count("---\n"), 2, text)
        self.assertIn("title: 报告\n", text)
        self.assertIn("tags: [a, b]\n", text)
        self.assertEqual(len(re.findall(r"^AIGC:", text, re.M)), 1, text)
        self.assertTrue(text.endswith("---\n# 正文\n"), text)
        self.assertEqual(aite_label.read_label(p), f2)

    # ------------------------------------------------------------ PDF

    @unittest.skipUnless(HAVE_PYPDF, "镜像里没有 pypdf（D12 未批）")
    def test_pdf_info_roundtrip(self):
        from pypdf import PdfReader, PdfWriter

        p = self.path("a.pdf")
        w = PdfWriter()
        w.add_blank_page(72, 72)
        w.add_metadata({"/Title": "月报"})
        with open(p, "wb") as f:
            w.write(f)
        res = aite_label.label_file(p, _fields())
        self.assertEqual(res["status"], "labeled", res)
        self.assertEqual(aite_label.read_label(p), _fields())
        meta = PdfReader(p).metadata
        self.assertEqual(meta.get("/Title"), "月报")
        self.assertEqual(len(PdfReader(p).pages), 1)

    def test_pdf_skipped_when_pypdf_missing(self):
        p = self.path("a.pdf")
        raw = b"%PDF-1.4\n%fake\n"
        with open(p, "wb") as f:
            f.write(raw)
        with mock.patch.dict(sys.modules, {"pypdf": None}):
            res = aite_label.label_file(p, _fields())
        self.assertEqual(res["status"], "skipped", res)
        self.assertEqual(res["reason"], aite_label.PDF_SKIP_REASON)
        self.assertEqual(aite_label.exit_code_for([res]), 3)
        with open(p, "rb") as f:
            self.assertEqual(f.read(), raw)

    # ------------------------------------------------------------ CLI / 载荷

    def test_unsupported_type_skipped_exit_3(self):
        p = self.path("a.csv")
        with open(p, "w", encoding="utf-8") as f:
            f.write("a,b\n1,2\n")
        proc = subprocess.run([sys.executable, LABELER, p], capture_output=True, text=True)
        self.assertEqual(proc.returncode, 3, proc.stderr)
        line = json.loads(proc.stdout)
        self.assertEqual(line["status"], "skipped")
        self.assertEqual(line["format"], "csv")
        with open(p, encoding="utf-8") as f:
            self.assertEqual(f.read(), "a,b\n1,2\n")

    def test_cli_prints_one_json_line_per_path(self):
        png = self.make_png()
        md = self.path("b.md")
        with open(md, "w", encoding="utf-8") as f:
            f.write("hi\n")
        csv = self.path("c.csv")
        with open(csv, "w", encoding="utf-8") as f:
            f.write("x\n")
        missing = self.path("nope.png")
        proc = subprocess.run(
            [sys.executable, LABELER, "--propagator", "飞书", png, md, csv],
            capture_output=True,
            text=True,
        )
        self.assertEqual(proc.returncode, 3, proc.stderr)
        lines = [json.loads(s) for s in proc.stdout.splitlines()]
        self.assertEqual([x["path"] for x in lines], [png, md, csv])
        for x in lines:
            self.assertEqual(set(x), {"path", "format", "status", "reason"})
        self.assertEqual([x["status"] for x in lines], ["labeled", "labeled", "skipped"])
        a, b = aite_label.read_label(png), aite_label.read_label(md)
        self.assertEqual(a["ProduceID"], b["ProduceID"])
        self.assertEqual(a["ContentPropagator"], "飞书")
        proc = subprocess.run([sys.executable, LABELER, missing], capture_output=True, text=True)
        self.assertEqual(proc.returncode, 1)
        self.assertEqual(json.loads(proc.stdout)["status"], "error")
        proc = subprocess.run([sys.executable, LABELER], capture_output=True, text=True)
        self.assertEqual(proc.returncode, 2)

    def test_payload_has_exactly_five_fields(self):
        p = self.make_png()
        aite_label.label_file(p)
        got = aite_label.read_label(p)
        self.assertEqual(set(got), FIVE)
        self.assertTrue(all(isinstance(v, str) for v in got.values()), got)
        self.assertEqual(got["Label"], "1")
        self.assertEqual(got["ContentProducer"], "Aite")
        self.assertRegex(got["ProduceID"], r"^[0-9a-f]{32}$")
        self.assertEqual((got["ContentPropagator"], got["PropagateID"]), ("", ""))
        self.assertEqual(set(aite_label.make_fields()), FIVE)

    # ------------------------------------------------------------ 可见水印

    def test_watermark_off_by_default(self):
        png, jpg = self.make_png(), self.make_jpeg()
        before = (_pixels(png), _pixels(jpg))
        self.assertEqual(aite_label.label_file(png)["status"], "labeled")
        self.assertEqual(aite_label.label_file(jpg)["status"], "labeled")
        proc = subprocess.run([sys.executable, LABELER, png, jpg], capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        self.assertEqual((_pixels(png), _pixels(jpg)), before)

    def test_watermark_opt_in_changes_pixels(self):
        png = self.make_png()
        before = _pixels(png)
        res = aite_label.label_file(png, _fields(), watermark=True)
        self.assertEqual(res["status"], "labeled", res)
        after = _pixels(png)
        self.assertEqual(after[1], before[1])
        self.assertNotEqual(after[2], before[2])
        self.assertEqual(aite_label.read_label(png), _fields())


if __name__ == "__main__":
    unittest.main()
